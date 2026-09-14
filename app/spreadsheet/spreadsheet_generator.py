from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.settings import get_settings


class SpreadsheetGenerator:
    def __init__(self) -> None:
        self.output_dir = Path(get_settings().reports_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def create_spreadsheet(
        self,
        city_name: str,
        city_key: str,
        hourly_data: list[dict],
        filters: dict[str, str | None],
        report_name: str,
    ) -> str:
        file_path = self.output_dir / report_name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Weather report"

        headers = [
            "City name",
            "City key",
            "Time",
            "Temperature (°C)",
            "Humidity (%)",
            "Wind speed (km/h)",
            "Weather code",
            "Filter date",
            "Filter period",
        ]
        worksheet.append(headers)

        for hourly_weather in hourly_data:
            worksheet.append(
                [
                    city_name,
                    city_key,
                    hourly_weather["time"],
                    hourly_weather["temperature"],
                    hourly_weather["humidity"],
                    hourly_weather["wind_speed"],
                    hourly_weather["weather_code"],
                    filters.get("day") or "All",
                    filters.get("period") or "All",
                ]
            )

        header_fill = PatternFill("solid", fgColor="075985")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            longest_value = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                longest_value + 2,
                32,
            )

        workbook.save(file_path)
        return str(file_path)


spreadsheet_generator = SpreadsheetGenerator()
